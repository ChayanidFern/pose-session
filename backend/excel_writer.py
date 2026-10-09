# ═══════════════════════════════════════════════════════
#  Neck GT Controller — Excel Writer
#  สร้างไฟล์ .xlsx ด้วย openpyxl
#  แก้ไขสไตล์และคอลัมน์ที่นี่
# ═══════════════════════════════════════════════════════

import os
from datetime import datetime
from typing import List, Dict, Any

import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter

# ── สีธีม ──────────────────────────────────────────────
COLORS = {
    "header_bg":  "1C2333",   # พื้นหลัง header เข้ม
    "header_fg":  "E6EDF3",   # ตัวอักษร header สว่าง
    "accent":     "3FB950",   # สีเขียว accent
    "accent2":    "79C0FF",   # สีฟ้า accent
    "row_odd":    "F6F8FA",   # แถวคี่
    "row_even":   "FFFFFF",   # แถวคู่
    "risk_normal":"EAF3DE",   # พื้นหลัง Normal risk
    "risk_low":   "FAEEDA",   # พื้นหลัง Low risk
    "risk_high":  "FCEBEB",   # พื้นหลัง High risk
}

# ── Activity → Risk Level mapping ──────────────────────
ACT_RISK = {
    "REST":    ("NORMAL",   COLORS["risk_normal"]),
    "PC":      ("LOW",      COLORS["risk_low"]),
    "NOTE":    ("LOW-MID",  COLORS["risk_low"]),
    "BOOK":    ("LOW-MID",  COLORS["risk_low"]),
    "TABLET":  ("MID",      COLORS["risk_low"]),
    "PHONE_H": ("HIGH",     COLORS["risk_high"]),
    "PHONE_D": ("HIGH",     COLORS["risk_high"]),
    "MIXED":   ("NATURAL",  COLORS["row_odd"]),
}

ACT_NAME_TH = {
    "REST":    "พัก / นั่งตรง",
    "NOTE":    "จดโน้ต",
    "TABLET":  "ไอแพด",
    "PHONE_H": "โทรศัพท์ (มือ)",
    "PHONE_D": "โทรศัพท์ (โต๊ะ)",
    "PC":      "คอมพิวเตอร์",
    "BOOK":    "หนังสือ",
    "MIXED":   "Naturalistic",
}


def _border(style="thin"):
    s = Side(style=style, color="D0D7DE")
    return Border(left=s, right=s, top=s, bottom=s)

def _header_font():
    return Font(name="Sarabun", bold=True, color=COLORS["header_fg"], size=10)

def _body_font(bold=False):
    return Font(name="Sarabun", bold=bold, size=10)

def _header_fill():
    return PatternFill("solid", fgColor=COLORS["header_bg"])

def _center():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)

def _left():
    return Alignment(horizontal="left", vertical="center", wrap_text=True)


# ── Sheet 1: ข้อมูล GT ────────────────────────────────
def _write_gt_sheet(ws, steps: List[Dict], sess: Dict):
    ws.title = "ข้อมูล GT"

    # Headers
    headers = [
        ("ลำดับ",            6),
        ("กิจกรรม",          12),
        ("ชื่อกิจกรรม",      18),
        ("Risk Level",        12),
        ("คำสั่ง",            45),
        ("เวลาเริ่ม (ms)",   14),
        ("เวลาสิ้นสุด (ms)", 16),
        ("ระยะเวลา (ms)",    14),
        ("ระยะเวลา (s)",     12),
        ("Participant",       12),
        ("Rater",             12),
        ("Session",           20),
        ("Scenario",          10),
        ("Split",             8),
        ("Timestamp",         22),
    ]

    # Title row
    ws.merge_cells("A1:O1")
    title_cell = ws["A1"]
    title_cell.value = f"Neck GT Data — {sess['sessId']} | {sess['sc']} | {sess['pid']}"
    title_cell.font  = Font(name="Sarabun", bold=True, size=13, color=COLORS["accent"])
    title_cell.fill  = PatternFill("solid", fgColor=COLORS["header_bg"])
    title_cell.alignment = _center()
    ws.row_dimensions[1].height = 28

    # Header row
    for col_idx, (h, w) in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col_idx, value=h)
        cell.font      = _header_font()
        cell.fill      = _header_fill()
        cell.alignment = _center()
        cell.border    = _border()
        ws.column_dimensions[get_column_letter(col_idx)].width = w
    ws.row_dimensions[2].height = 22

    # Data rows
    ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    for i, step in enumerate(steps):
        row = i + 3
        act   = step.get("act", "")
        risk_label, risk_color = ACT_RISK.get(act, ("—", COLORS["row_odd"]))
        bg    = risk_color if i % 2 == 0 else COLORS["row_even"]
        fill  = PatternFill("solid", fgColor=bg)

        values = [
            i + 1,
            act,
            ACT_NAME_TH.get(act, act),
            risk_label,
            step.get("th", ""),
            step.get("onsetMs", 0),
            step.get("offsetMs", 0),
            step.get("durationMs", 0),
            round(step.get("durationMs", 0) / 1000, 2),
            sess.get("pid", ""),
            sess.get("rater", ""),
            sess.get("sessId", ""),
            sess.get("sc", ""),
            sess.get("split", ""),
            ts,
        ]

        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col_idx, value=val)
            cell.fill      = fill
            cell.border    = _border()
            cell.font      = _body_font()
            cell.alignment = _center() if col_idx != 5 else _left()

        ws.row_dimensions[row].height = 18

    # Freeze header rows
    ws.freeze_panes = "A3"


# ── Sheet 2: สรุปกิจกรรม ──────────────────────────────
def _write_summary_sheet(ws, steps: List[Dict]):
    ws.title = "สรุปกิจกรรม"

    headers = [
        ("กิจกรรม",      14),
        ("ชื่อกิจกรรม",  20),
        ("Risk Level",    14),
        ("จำนวน Step",   12),
        ("รวม (ms)",      14),
        ("รวม (s)",       10),
        ("รวม (นาที)",    12),
        ("% ของทั้งหมด",  14),
    ]

    # Title
    ws.merge_cells("A1:H1")
    t = ws["A1"]
    t.value = "สรุปรายกิจกรรม"
    t.font  = Font(name="Sarabun", bold=True, size=13, color=COLORS["accent"])
    t.fill  = PatternFill("solid", fgColor=COLORS["header_bg"])
    t.alignment = _center()
    ws.row_dimensions[1].height = 28

    for col_idx, (h, w) in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col_idx, value=h)
        cell.font      = _header_font()
        cell.fill      = _header_fill()
        cell.alignment = _center()
        cell.border    = _border()
        ws.column_dimensions[get_column_letter(col_idx)].width = w
    ws.row_dimensions[2].height = 22

    # Aggregate
    summary: Dict[str, Dict] = {}
    total_ms = 0
    for step in steps:
        act = step.get("act", "")
        dur = step.get("durationMs", 0)
        total_ms += dur
        if act not in summary:
            summary[act] = {"count": 0, "total_ms": 0}
        summary[act]["count"]    += 1
        summary[act]["total_ms"] += dur

    for row_idx, (act, data) in enumerate(summary.items(), start=3):
        risk_label, risk_color = ACT_RISK.get(act, ("—", COLORS["row_odd"]))
        fill = PatternFill("solid", fgColor=risk_color)
        pct  = round(data["total_ms"] / total_ms * 100, 1) if total_ms else 0

        values = [
            act,
            ACT_NAME_TH.get(act, act),
            risk_label,
            data["count"],
            data["total_ms"],
            round(data["total_ms"] / 1000, 2),
            round(data["total_ms"] / 60000, 2),
            f"{pct}%",
        ]
        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill      = fill
            cell.border    = _border()
            cell.font      = _body_font()
            cell.alignment = _center()
        ws.row_dimensions[row_idx].height = 18

    ws.freeze_panes = "A3"


# ── Sheet 3: Metadata ──────────────────────────────────
def _write_metadata_sheet(ws, sess: Dict, step_count: int):
    ws.title = "Metadata"

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 36

    # Title
    ws.merge_cells("A1:B1")
    t = ws["A1"]
    t.value = "Session Metadata"
    t.font  = Font(name="Sarabun", bold=True, size=13, color=COLORS["accent"])
    t.fill  = PatternFill("solid", fgColor=COLORS["header_bg"])
    t.alignment = _center()
    ws.row_dimensions[1].height = 28

    rows = [
        ("Session ID",      sess.get("sessId",  "—")),
        ("Scenario",        sess.get("sc",       "—")),
        ("Participant ID",  sess.get("pid",      "—")),
        ("Split",           sess.get("split",    "—")),
        ("Rater",           sess.get("rater",    "—")),
        ("Total Steps",     step_count),
        ("Note",            sess.get("note",     "")),
        ("Export Time",     datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]

    for row_idx, (field, value) in enumerate(rows, start=2):
        bg = COLORS["row_odd"] if row_idx % 2 == 0 else COLORS["row_even"]

        k = ws.cell(row=row_idx, column=1, value=field)
        k.font      = _body_font(bold=True)
        k.fill      = PatternFill("solid", fgColor=bg)
        k.border    = _border()
        k.alignment = _left()

        v = ws.cell(row=row_idx, column=2, value=value)
        v.font      = _body_font()
        v.fill      = PatternFill("solid", fgColor=bg)
        v.border    = _border()
        v.alignment = _left()

        ws.row_dimensions[row_idx].height = 18


# ── Master file: append ทุก session ต่อกัน ─────────────
MASTER_FILE = "all_sessions.xlsx"

MASTER_HEADERS = [
    ("ลำดับ",             6),
    ("Session ID",        22),
    ("Participant",       12),
    ("Scenario",          10),
    ("Split",              8),
    ("Rater",             12),
    ("กิจกรรม",           12),
    ("ชื่อกิจกรรม",       18),
    ("Risk Level",        12),
    ("คำสั่ง",             44),
    ("เวลาเริ่ม (ms)",    14),
    ("เวลาสิ้นสุด (ms)", 16),
    ("ระยะเวลา (ms)",    14),
    ("ระยะเวลา (s)",     12),
    ("วันที่เริ่ม Session", 22),
    ("วันที่จบ Session",   22),
    ("สถานะ",             10),
    ("Note",              24),
    ("บันทึกเมื่อ",        22),
]

def _parse_times(note_raw: str):
    """แยก start/end time และ note จริงออกจาก note field"""
    import re
    is_partial = "[PARTIAL]" in note_raw
    sm = re.search(r'start:([\d/: ]+)', note_raw)
    em = re.search(r'end:([\d/: ]+?)(?:\s*\[|$)', note_raw)
    start_time = sm.group(1).strip() if sm else ""
    end_time   = em.group(1).strip() if em else ""
    note_clean = re.sub(r'\|.*', '', note_raw).strip()
    return start_time, end_time, is_partial, note_clean


def save_session_to_excel(sess: Dict, data_dir: str) -> str:
    """
    Append ข้อมูล session ต่อท้าย all_sessions.xlsx (ไฟล์เดียว)
    ถ้าหยุดกลางคันก็บันทึกได้ — มี timestamp แม่นยำ
    """
    steps    = sess.get("steps", [])
    sess_id  = sess.get("sessId", f"sess_{int(datetime.now().timestamp())}")
    sc       = sess.get("sc", "SC-00")
    pid      = sess.get("pid", "P00")
    split    = sess.get("split", "train")
    rater    = sess.get("rater", "")
    note_raw = sess.get("note", "")

    # แยก timestamp จาก note
    start_time, end_time, is_partial, note_clean = _parse_times(note_raw)
    status   = "PARTIAL" if is_partial else "COMPLETE"

    # เวลาบันทึกจริง (server time UTC+7)
    from datetime import timezone, timedelta
    bkk = timezone(timedelta(hours=7))
    now_bkk = datetime.now(bkk).strftime("%d/%m/%Y %H:%M:%S")
    if not start_time: start_time = now_bkk
    if not end_time:   end_time   = now_bkk

    master_path = os.path.join(data_dir, MASTER_FILE)

    # โหลดไฟล์เดิม หรือสร้างใหม่
    if os.path.exists(master_path):
        wb = openpyxl.load_workbook(master_path)
        ws = wb["ข้อมูล GT"]
        next_row = ws.max_row + 1
        row_count = ws.max_row - 1  # ไม่นับ header
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "ข้อมูล GT"
        # เขียน header ครั้งแรกครั้งเดียว
        for col_idx, (h, w) in enumerate(MASTER_HEADERS, 1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.font      = _header_font()
            cell.fill      = _header_fill()
            cell.alignment = _center()
            cell.border    = _border()
            ws.column_dimensions[get_column_letter(col_idx)].width = w
        ws.row_dimensions[1].height = 22
        ws.freeze_panes = "A2"
        next_row  = 2
        row_count = 0

    global_idx = row_count + 1

    for step in steps:
        act   = step.get("act", "")
        risk_label, risk_color = ACT_RISK.get(act, ("—", COLORS["row_even"]))
        bg    = risk_color if global_idx % 2 == 1 else COLORS["row_even"]
        fill  = PatternFill("solid", fgColor=bg)

        values = [
            global_idx,
            sess_id, pid, sc, split, rater,
            act,
            ACT_NAME_TH.get(act, act),
            risk_label,
            step.get("th", ""),
            step.get("onsetMs", 0),
            step.get("offsetMs", 0),
            step.get("durationMs", 0),
            round(step.get("durationMs", 0) / 1000, 2),
            start_time,
            end_time,
            status,
            note_clean,
            now_bkk,
        ]

        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=next_row, column=col_idx, value=val)
            cell.fill      = fill
            cell.font      = _body_font()
            cell.border    = _border()
            cell.alignment = _left() if col_idx == 10 else _center()
        ws.row_dimensions[next_row].height = 18

        next_row   += 1
        global_idx += 1

    wb.save(master_path)
    return MASTER_FILE


# ── List sessions ─────────────────────────────────────
def list_sessions(data_dir: str) -> list:
    """
    แสดงรายการ session ทั้งหมดใน /data/
    """
    files = []
    for f in sorted(os.listdir(data_dir), reverse=True):
        if f.endswith(".xlsx"):
            path = os.path.join(data_dir, f)
            stat = os.stat(path)
            files.append({
                "filename": f,
                "size_kb":  round(stat.st_size / 1024, 1),
                "created":  datetime.fromtimestamp(stat.st_ctime).strftime("%Y-%m-%d %H:%M"),
                "download": f"/api/download/{f}",
            })
    return files
